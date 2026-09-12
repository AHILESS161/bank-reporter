import csv
import html
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Artifact, DocumentVersion, FinancialFact, ProvenanceRef, Report, SourceDocument
from .browser import BrowserClient
from .financial import growth
from .lieflat import LIEFLAT_COMMIT, format_fact_value, report_contract, rung_chart, select_chart
from .model_router import ModelRouter, ModelUnavailable
from .security import file_sha256
from .storage import Storage


class ReportService:
    def __init__(self, db: Session, analysis_run_id: str | None = None):
        self.db = db
        self.storage = Storage()
        self.models = ModelRouter(db, analysis_run_id)

    def create(self, report: Report, question: str, output_formats: list[str]) -> Report:
        report.status = "processing"
        self.db.commit()
        rows = self._facts(report.document_ids)
        try:
            narrative = self.models.finance_narrative(rows, question)
        except ModelUnavailable as exc:
            narrative = {
                "summary": "Финансовая интерпретация не сформирована.",
                "highlights": [],
                "risks": [str(exc)],
                "chart": {"template": "F1", "metric_codes": []},
            }
            report.status = "partial"
        report.summary = narrative["summary"]
        html_path = self.storage.artifact_path(report.id, "report.html")
        html_path.write_text(self._html(report, narrative, rows), encoding="utf-8")
        self._artifact(report, "html", "text/html", html_path)
        if "xlsx" in output_formats:
            xlsx = self.storage.artifact_path(report.id, "report.xlsx")
            self._xlsx(xlsx, narrative, rows, self._calculations(rows))
            self._artifact(
                report, "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx
            )
        if "csv" in output_formats:
            csv_path = self.storage.artifact_path(report.id, "facts.csv")
            self._csv(csv_path, rows)
            self._artifact(report, "csv", "text/csv", csv_path)
        if "pdf" in output_formats or "png" in output_formats:
            pdf = self.storage.artifact_path(report.id, "report.pdf") if "pdf" in output_formats else None
            png = self.storage.artifact_path(report.id, "report.png") if "png" in output_formats else None
            try:
                BrowserClient().render(str(html_path), str(pdf) if pdf else None, str(png) if png else None)
                if pdf and pdf.exists():
                    self._artifact(report, "pdf", "application/pdf", pdf)
                if png and png.exists():
                    self._artifact(report, "png", "image/png", png)
            except Exception:
                report.status = "partial"
        if report.status != "partial":
            report.status = "completed"
        self.db.commit()
        self.db.refresh(report)
        return report

    def _facts(self, document_ids: list[str]) -> list[dict]:
        statement = (
            select(FinancialFact, ProvenanceRef, SourceDocument)
            .join(ProvenanceRef, FinancialFact.provenance_id == ProvenanceRef.id)
            .join(DocumentVersion, FinancialFact.document_version_id == DocumentVersion.id)
            .join(SourceDocument, DocumentVersion.document_id == SourceDocument.id)
            .where(SourceDocument.id.in_(document_ids))
            .order_by(FinancialFact.period_end, FinancialFact.metric_code)
        )
        return [
            {
                "id": fact.id,
                "metric_code": fact.metric_code,
                "label": fact.label,
                "value": str(fact.value),
                "currency": fact.currency,
                "unit_scale": fact.unit_scale,
                "period_end": fact.period_end,
                "confidence": fact.confidence,
                "document": document.title,
                "source_url": provenance.source_url,
                "page": provenance.page,
                "sheet": provenance.sheet,
                "cell_range": provenance.cell_range,
            }
            for fact, provenance, document in self.db.execute(statement).all()
        ]

    def _html(self, report: Report, narrative: dict, facts: list[dict]) -> str:
        contract = report_contract(report.report_kind, self.models.settings.lieflat_dir)
        requested_chart = narrative.get("chart", {}).get("template", "F1")
        chart_id = select_chart(requested_chart, facts)
        chart = rung_chart(facts, chart_id)
        fact_by_id = {str(item["id"]): item for item in facts}
        source_number = {str(item["id"]): index for index, item in enumerate(facts, 1)}

        def present_text(text: str, fact_ids: list[str]) -> str:
            rendered = text
            for fact_id in fact_ids:
                fact = fact_by_id.get(str(fact_id))
                if fact:
                    rendered = rendered.replace(str(fact["value"]), format_fact_value(fact))
            return rendered

        def coordinate(item: dict) -> str:
            if item.get("page"):
                return f"страница {item['page']}"
            if item.get("sheet"):
                return f"лист {item['sheet']}"
            if item.get("cell_range"):
                return f"ячейка {item['cell_range']}"
            return "координата источника не указана"

        highlights = "".join(
            f'<li data-provenance="{html.escape(",".join(item.get("fact_ids", [])))}">'
            f'<span>{html.escape(present_text(item.get("text", ""), item.get("fact_ids", [])))}</span>'
            + "".join(
                f'<a class="footnote" href="#source-{html.escape(str(fact_id))}">[{source_number.get(str(fact_id), "")}]</a>'
                for fact_id in item.get("fact_ids", [])
                if str(fact_id) in source_number
            )
            + "</li>"
            for item in narrative.get("highlights", [])
        )
        risks = "".join(f"<li>{html.escape(str(item))}</li>" for item in narrative.get("risks", []))
        sources = "".join(
            f'<li id="source-{html.escape(str(item["id"]))}"><span class="source-no">{index:02d}</span><div>'
            f'<a href="{html.escape(str(item.get("source_url") or "#"))}" target="_blank" rel="noreferrer">{html.escape(item["document"])}</a>'
            f'<small>{html.escape(coordinate(item))}</small>'
            f'</div><code>{html.escape(str(item["id"]))}</code></li>'
            for index, item in enumerate(facts, 1)
        )
        kpis = "".join(
            f'<article><span>{html.escape(str(item.get("label") or item.get("metric_code")))}</span>'
            f'<strong>{html.escape(format_fact_value(item))}</strong>'
            f'<small>{html.escape(str(item.get("period_end") or "Период не указан"))}</small></article>'
            for item in facts[:4]
        )
        width = "600px" if contract.report_id == "R11" else "1080px"
        return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{html.escape(report.title)}</title>
<!-- Lieflat source: {contract.source_file} @ {LIEFLAT_COMMIT}; candidates: {", ".join(contract.candidates)}; {html.escape(contract.selection_reason)} -->
<style>
:root{{--page:#f4f1eb;--card:#fffdfa;--ink:#17211d;--muted:#66706b;--line:#dce0dc;--accent:#176b5b;--accent-soft:#dfeee9;--track:#e9ece9;--danger:#a84c43}}
*{{box-sizing:border-box}}html{{background:var(--page)}}body{{margin:0;background:var(--page);color:var(--ink);font-family:Inter,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif;font-variant-numeric:tabular-nums;padding:34px 20px}}a{{color:var(--accent)}}.sheet{{width:{width};max-width:100%;margin:auto;background:var(--card);border:1px solid var(--line);border-radius:24px;box-shadow:0 18px 55px rgba(23,33,29,.07);overflow:hidden}}.hero{{padding:44px 48px 38px;background:linear-gradient(135deg,#fdfcf9 0%,#edf5f1 100%);border-bottom:1px solid var(--line)}}.hero-top{{display:flex;align-items:center;justify-content:space-between;gap:20px}}.eyebrow,.section-label{{margin:0;font-size:10px;font-weight:800;letter-spacing:.15em;text-transform:uppercase;color:var(--accent)}}.verified{{padding:7px 10px;border:1px solid #b8d7cd;border-radius:99px;background:#eef8f4;color:var(--accent);font-size:9px;font-weight:800;text-transform:uppercase;letter-spacing:.08em}}h1{{max-width:780px;margin:28px 0 18px;font-size:48px;line-height:1.02;letter-spacing:-.045em}}.lead{{max-width:850px;margin:0;color:#3e4a45;font-size:18px;line-height:1.55}}.content{{padding:34px 48px 42px}}.kpis{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px;margin-bottom:34px}}.kpis article{{min-width:0;padding:17px 18px;border:1px solid var(--line);border-radius:14px;background:#fbfbf8}}.kpis span,.kpis small{{display:block;color:var(--muted);font-size:10px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}.kpis strong{{display:block;margin:12px 0 8px;font-size:19px;letter-spacing:-.02em;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}.analysis-grid{{display:grid;grid-template-columns:minmax(0,1.65fr) minmax(270px,.8fr);gap:24px;align-items:start}}.chart-card,.insight-card{{border:1px solid var(--line);border-radius:18px;background:#fff;padding:24px}}.chart-card h2,.insight-card h2{{margin:7px 0 22px;font-size:21px;letter-spacing:-.025em}}.lieflat-chart{{display:block;width:100%;height:auto}}.insight-card ol,.insight-card ul{{margin:0;padding-left:20px}}.insight-card li{{padding:0 0 14px 3px;line-height:1.5;font-size:14px}}.insight-card li:last-child{{padding-bottom:0}}.limitations{{margin-top:24px;padding-top:22px;border-top:1px solid var(--line)}}.footnote{{margin-left:4px;text-decoration:none;font-size:9px;vertical-align:super}}.sources{{margin-top:28px;border:1px solid var(--line);border-radius:16px;background:#f8f8f5;overflow:hidden}}.sources summary{{display:flex;justify-content:space-between;align-items:center;gap:16px;padding:17px 20px;cursor:pointer;list-style:none;font-size:13px;font-weight:800}}.sources summary::-webkit-details-marker{{display:none}}.sources summary:after{{content:"＋";font-size:17px;color:var(--accent)}}.sources[open] summary:after{{content:"−"}}.sources summary strong{{margin-left:auto;color:var(--muted);font-size:10px;font-weight:600}}.source-list{{margin:0;padding:0 20px 14px;list-style:none;border-top:1px solid var(--line)}}.source-list li{{display:grid;grid-template-columns:30px minmax(0,1fr) auto;gap:12px;align-items:start;padding:14px 0;border-bottom:1px solid var(--line)}}.source-list li:last-child{{border-bottom:0}}.source-no{{color:var(--accent);font-size:10px;font-weight:800}}.source-list a{{font-size:12px;font-weight:700}}.source-list small{{display:block;margin-top:5px;color:var(--muted);font-size:10px}}.source-list code{{max-width:150px;color:#88908c;font-size:8px;overflow:hidden;text-overflow:ellipsis}}footer{{display:flex;justify-content:space-between;gap:20px;margin-top:28px;padding-top:18px;border-top:1px solid var(--line);color:var(--muted);font-size:9px;line-height:1.5}}footer b{{color:var(--ink)}}
@media(max-width:780px){{body{{padding:0}}.sheet{{border:0;border-radius:0}}.hero,.content{{padding:28px 22px}}h1{{font-size:36px}}.kpis{{grid-template-columns:1fr 1fr}}.analysis-grid{{grid-template-columns:1fr}}.source-list li{{grid-template-columns:25px 1fr}}.source-list code{{display:none}}}}
@media print{{@page{{size:A4;margin:12mm}}body{{padding:0;background:white}}.sheet{{width:100%;border:0;border-radius:0;box-shadow:none}}.hero{{padding:12mm 10mm 9mm}}.content{{padding:8mm 10mm}}h1{{font-size:34px}}.sources{{break-before:page}}.sources>summary{{display:none}}.sources>.source-list{{display:block!important}}}}
</style></head>
<body><main class="sheet" data-lieflat-report="{contract.report_id}" data-lieflat-chart="{chart_id}" data-lieflat-commit="{LIEFLAT_COMMIT}">
<header class="hero"><div class="hero-top"><p class="eyebrow">BANK REPORTER · {contract.report_id}</p><span class="verified">Проверяемый анализ</span></div><h1>{html.escape(report.title)}</h1><p class="lead">{html.escape(narrative["summary"])}</p></header>
<div class="content">{f'<section class="kpis">{kpis}</section>' if kpis else ''}<div class="analysis-grid"><section class="chart-card"><p class="section-label">Финансовый профиль</p><h2>Ключевые показатели</h2>{chart if facts else '<p>Нормализованные показатели не найдены.</p>'}</section><aside class="insight-card"><p class="section-label">Редакционный вывод</p><h2>Что важно</h2><ol>{highlights or '<li>Требуется ручная проверка извлечения.</li>'}</ol><div class="limitations"><p class="section-label">Ограничения</p><ul>{risks or '<li>Существенные ограничения не указаны.</li>'}</ul></div></aside></div>
<details class="sources"><summary><span>Источники и координаты</span><strong>{len(facts)} записей</strong></summary><ol class="source-list">{sources or '<li>Источники не привязаны.</li>'}</ol></details>
<footer><span><b>Bank Reporter</b><br>Не является инвестиционной рекомендацией.</span><span>Каждое число связано с ProvenanceRef.<br>Lieflat Charts · PolyForm Noncommercial 1.0.0.</span></footer></div></main></body></html>'''

    @staticmethod
    def _xlsx(path: Path, narrative: dict, facts: list[dict], calculations: list[dict]) -> None:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Выводы"
        summary.append(["Резюме", narrative.get("summary")])
        summary.append([])
        for item in narrative.get("highlights", []):
            summary.append([item.get("text"), ",".join(item.get("fact_ids", []))])
        facts_sheet = workbook.create_sheet("Факты")
        facts_sheet.append(
            ["ID", "Показатель", "Название", "Значение", "Валюта", "Масштаб", "Период", "Уверенность"]
        )
        for item in facts:
            facts_sheet.append(
                [
                    item["id"],
                    item["metric_code"],
                    item["label"],
                    item["value"],
                    item["currency"],
                    item["unit_scale"],
                    str(item["period_end"] or ""),
                    item["confidence"],
                ]
            )
        calculation_sheet = workbook.create_sheet("Расчеты")
        calculation_sheet.append(
            ["Показатель", "Предыдущий период", "Текущий период", "Изменение, %", "Fact IDs"]
        )
        for item in calculations:
            calculation_sheet.append(
                [
                    item["metric_code"],
                    item["previous_period"],
                    item["current_period"],
                    item["growth_percent"],
                    ",".join(item["fact_ids"]),
                ]
            )
        sources = workbook.create_sheet("Источники")
        sources.append(["Документ", "URL", "Страница", "Лист", "Ячейка"])
        for item in facts:
            sources.append(
                [item["document"], item["source_url"], item["page"], item["sheet"], item["cell_range"]]
            )
        workbook.save(path)

    @staticmethod
    def _calculations(facts: list[dict]) -> list[dict]:
        grouped: dict[str, list[dict]] = {}
        for item in facts:
            if item.get("period_end"):
                grouped.setdefault(str(item["metric_code"]), []).append(item)
        output: list[dict] = []
        for metric_code, items in grouped.items():
            ordered = sorted(items, key=lambda item: item["period_end"])
            for previous, current in zip(ordered, ordered[1:], strict=False):
                change = growth(Decimal(str(current["value"])), Decimal(str(previous["value"])))
                output.append(
                    {
                        "metric_code": metric_code,
                        "previous_period": str(previous["period_end"]),
                        "current_period": str(current["period_end"]),
                        "growth_percent": str(change) if change is not None else "n/a",
                        "fact_ids": [str(previous["id"]), str(current["id"])],
                    }
                )
        return output

    @staticmethod
    def _csv(path: Path, facts: list[dict]) -> None:
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(
                stream, fieldnames=list(facts[0].keys()) if facts else ["metric_code", "value"]
            )
            writer.writeheader()
            writer.writerows(facts)

    def _artifact(self, report: Report, fmt: str, mime: str, path: Path) -> None:
        artifact = self.db.scalar(
            select(Artifact).where(
                Artifact.report_id == report.id,
                Artifact.format == fmt,
            )
        )
        if artifact is None:
            artifact = Artifact(
                report_id=report.id,
                format=fmt,
            )
            self.db.add(artifact)
        artifact.mime_type = mime
        artifact.storage_path = str(path)
        artifact.size_bytes = path.stat().st_size
        artifact.sha256 = file_sha256(path)
        self.db.flush()
