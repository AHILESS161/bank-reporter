import csv
import io
import re
import zipfile
from xml.etree import ElementTree
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import fitz
import openpyxl
import pdfplumber
import pytesseract
from charset_normalizer import from_bytes
from dbfread import DBF
from PIL import Image

from ..config import get_settings


@dataclass
class CellFact:
    label: str
    value: str
    page: int | None = None
    sheet: str | None = None
    cell_range: str | None = None
    unit_scale: int = 1
    currency: str | None = "RUB"
    period_end: str | None = None


@dataclass
class ExtractionResult:
    text: str = ""
    tables: list[dict[str, Any]] = field(default_factory=list)
    candidates: list[CellFact] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def json(self) -> dict[str, Any]:
        return {
            "tables": self.tables,
            "candidates": [asdict(item) for item in self.candidates],
            "warnings": self.warnings,
        }


class Extractor:
    allowed = {".pdf", ".xlsx", ".csv", ".dbf", ".zip", ".xml"}

    def extract(self, path: Path) -> ExtractionResult:
        suffix = path.suffix.lower()
        if suffix not in self.allowed:
            raise ValueError(f"Неподдерживаемый формат: {suffix}")
        return {
            ".pdf": self._pdf,
            ".xlsx": self._xlsx,
            ".csv": self._csv,
            ".dbf": self._dbf,
            ".zip": self._zip,
            ".xml": self._xml,
        }[suffix](path)

    def _pdf(self, path: Path) -> ExtractionResult:
        result = ExtractionResult()
        doc = fitz.open(path)
        pages: list[str] = []
        page_texts: list[str] = []
        for number, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) < 30:
                pix = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                text = pytesseract.image_to_string(image, lang="rus+eng").strip()
                result.warnings.append(f"OCR использован для страницы {number}")
            pages.append(f"\n--- PAGE {number} ---\n{text}")
            page_texts.append(text)
        formal_document = len(page_texts) > 10 or any(
            marker in text.casefold()
            for text in page_texts[:3]
            for marker in ("финансовой отчетности", "финансовой отчётности", "пресс-релиз")
        )
        for number, text in enumerate(page_texts, start=1):
            structured = self._financial_statement_candidates(text, page=number)
            result.candidates.extend(structured)
            if not formal_document and not structured:
                result.candidates.extend(self._text_candidates(text, page=number))
        result.text = "".join(pages)
        try:
            with pdfplumber.open(path) as pdf:
                for number, page in enumerate(pdf.pages, start=1):
                    for index, table in enumerate(page.extract_tables() or []):
                        rows = [[str(cell or "").strip() for cell in row] for row in table]
                        result.tables.append(
                            {"name": f"page_{number}_table_{index + 1}", "page": number, "rows": rows}
                        )
                        if formal_document and not self._primary_financial_page(page_texts[number - 1]):
                            continue
                        result.candidates.extend(
                            self._pdf_table_candidates(rows, page_texts[number - 1], number)
                        )
        except Exception as exc:  # table extraction is best-effort
            result.warnings.append(f"Таблицы PDF не извлечены: {exc}")
        return result

    def _xlsx(self, path: Path) -> ExtractionResult:
        result = ExtractionResult()
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        chunks: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(values)
                    chunks.append(" | ".join(values))
                    for col, value in enumerate(values[1:], start=2):
                        if self._decimal(value) is not None and values[0]:
                            result.candidates.append(
                                CellFact(
                                    values[0],
                                    value,
                                    sheet=sheet.title,
                                    cell_range=f"{openpyxl.utils.get_column_letter(col)}{row_number}",
                                )
                            )
            result.tables.append({"name": sheet.title, "sheet": sheet.title, "rows": rows[:10_000]})
        result.text = "\n".join(chunks)
        return result

    def _csv(self, path: Path) -> ExtractionResult:
        raw = path.read_bytes()
        match = from_bytes(raw).best()
        text = str(match) if match else raw.decode("utf-8", errors="replace")
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        rows = list(csv.reader(io.StringIO(text), dialect))
        result = ExtractionResult(text=text, tables=[{"name": path.stem, "rows": rows[:50_000]}])
        for number, row in enumerate(rows, start=1):
            if len(row) >= 2 and self._decimal(row[1]) is not None:
                result.candidates.append(CellFact(row[0], row[1], sheet=path.name, cell_range=f"B{number}"))
        return result

    def _dbf(self, path: Path) -> ExtractionResult:
        table = DBF(path, load=True, char_decode_errors="replace")
        rows = [[str(record.get(name, "")) for name in table.field_names] for record in table]
        return ExtractionResult(
            text="\n".join(" | ".join(row) for row in rows),
            tables=[{"name": path.stem, "headers": table.field_names, "rows": rows}],
        )

    def _zip(self, path: Path) -> ExtractionResult:
        settings = get_settings()
        result = ExtractionResult()
        with zipfile.ZipFile(path) as archive:
            total = sum(item.file_size for item in archive.infolist())
            if total > settings.max_archive_mb * 1024 * 1024 or len(archive.infolist()) > 10_000:
                raise ValueError("Архив превышает безопасный лимит")
            for member in archive.infolist():
                member_name = Path(member.filename)
                if member.is_dir() or member_name.is_absolute() or ".." in member_name.parts:
                    continue
                if member_name.suffix.lower() not in self.allowed - {".zip"}:
                    continue
                target = get_settings().data_dir / "tmp" / f"extract-{path.stem}" / member_name.name
                target.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(member) as source, target.open("wb") as output:
                    output.write(source.read())
                nested = self.extract(target)
                result.text += nested.text
                result.tables.extend(nested.tables)
                result.candidates.extend(nested.candidates)
                result.warnings.extend(nested.warnings)
        return result

    def _xml(self, path: Path) -> ExtractionResult:
        root = ElementTree.fromstring(path.read_bytes())
        rows: list[list[str]] = []
        candidates: list[CellFact] = []
        for row_number, node in enumerate(root.iter(), start=1):
            children = list(node)
            if not children:
                continue
            values = {child.tag.rsplit("}", 1)[-1]: (child.text or "").strip() for child in children}
            if not any(values.values()):
                continue
            rows.append([f"{key}={value}" for key, value in values.items()])
            label = next(
                (value for key, value in values.items() if "name" in key.casefold() and value),
                node.tag.rsplit("}", 1)[-1],
            )
            for key, value in values.items():
                if self._decimal(value) is not None:
                    candidates.append(
                        CellFact(f"{label} {key}", value, sheet=path.name, cell_range=f"XML:{row_number}")
                    )
        text = "\n".join(" | ".join(row) for row in rows)
        return ExtractionResult(
            text=text,
            tables=[{"name": path.stem, "rows": rows[:50_000]}],
            candidates=candidates[:10_000],
        )

    def _text_candidates(self, text: str, page: int) -> list[CellFact]:
        candidates: list[CellFact] = []
        pattern = re.compile(r"^(.{3,100}?)\s+([()\-−]?\d[\d\s.,]*\)?)\s*$")
        for line in text.splitlines():
            match = pattern.match(line.strip())
            if match and self._decimal(match.group(2)) is not None:
                candidates.append(CellFact(match.group(1).strip(" ."), match.group(2), page=page))
        return candidates[:2000]

    def _financial_statement_candidates(self, text: str, page: int) -> list[CellFact]:
        """Read the common label / note / current / previous layout of bank PDF statements."""
        if not self._primary_financial_page(text):
            return []
        unit_scale, value_multiplier = self._financial_units(text)
        if unit_scale == 1:
            return []

        lines = [line.strip() for line in text.splitlines() if line.strip()]
        years = self._financial_years(text)
        if not years:
            return []

        aliases = tuple(
            alias
            for values in (
                ("всего активов", "совокупные активы", "итого активов", "total assets"),
                ("всего капитала", "собственные средства", "total equity"),
                ("чистая прибыль", "прибыль за год", "прибыль за период", "net profit"),
                ("чистый процентный доход", "net interest income"),
                ("чистые комиссионные доходы", "net fee"),
                ("операционные расходы",),
                ("кредиты клиентам", "loans to customers", "кредитный портфель"),
                ("средства клиентов", "customer accounts", "customer funds"),
                ("достаточность капитала", "н1.0", "capital adequacy"),
                ("рентабельность активов", "return on assets", "roa"),
                ("рентабельность капитала", "return on equity", "roe"),
            )
            for alias in values
        )
        output: list[CellFact] = []
        for index, label in enumerate(lines):
            label_lower = label.casefold()
            if not any(alias in label_lower for alias in aliases):
                continue
            values: list[str] = []
            for following in lines[index + 1 : index + 10]:
                number = self._decimal(following.rstrip("%"))
                if number is None:
                    if values or any(alias in following.casefold() for alias in aliases):
                        break
                    continue
                absolute = abs(number)
                if 1900 <= absolute <= 2100:
                    continue
                if not values and number == number.to_integral_value() and 0 <= absolute <= 40:
                    continue
                values.append(following)
                if len(values) >= len(years):
                    break
            for value_index, value in enumerate(values):
                normalized_value = value
                if not value.endswith("%") and value_multiplier != 1:
                    parsed_value = self._decimal(value)
                    normalized_value = str(parsed_value * value_multiplier) if parsed_value else value
                output.append(
                    CellFact(
                        label=label,
                        value=normalized_value,
                        page=page,
                        unit_scale=1 if value.endswith("%") else unit_scale,
                        currency="%" if value.endswith("%") else "RUB",
                        period_end=f"{years[min(value_index, len(years) - 1)]}-12-31",
                    )
                )
        return output

    def _pdf_table_candidates(
        self, rows: list[list[str]], page_text: str, page: int
    ) -> list[CellFact]:
        """Preserve table column years and page units instead of creating unitless facts."""
        years = self._financial_years(page_text)
        unit_scale, value_multiplier = self._financial_units(page_text)
        column_years: dict[int, int] = {}
        for header in rows[:6]:
            for column, cell in enumerate(header):
                match = re.search(r"\b(20\d{2})\b", cell)
                if match:
                    column_years[column] = int(match.group(1))

        output: list[CellFact] = []
        for row in rows:
            if len(row) < 2 or not row[0]:
                continue
            numeric = [
                (column, value)
                for column, value in enumerate(row[1:], start=1)
                if self._decimal(value.rstrip("%")) is not None
            ]
            row_years = dict(column_years)
            if not row_years and len(years) >= 2 and len(numeric) >= len(years):
                row_years = {
                    column: year
                    for (column, _), year in zip(numeric[-len(years) :], years, strict=True)
                }
            for column, value in numeric:
                if row_years and column not in row_years:
                    continue
                is_percent = value.endswith("%")
                normalized_value = value
                parsed_value = self._decimal(value.rstrip("%"))
                if not is_percent and value_multiplier != 1 and parsed_value is not None:
                    normalized_value = str(parsed_value * value_multiplier)
                year = row_years.get(column)
                output.append(
                    CellFact(
                        row[0],
                        normalized_value,
                        page=page,
                        unit_scale=1 if is_percent else unit_scale,
                        currency="%" if is_percent else "RUB",
                        period_end=f"{year}-12-31" if year else None,
                    )
                )
        return output

    @staticmethod
    def _financial_years(text: str) -> list[int]:
        years: list[int] = []
        for line in (line.strip() for line in text.splitlines()[:60]):
            for match in re.finditer(r"\b(20\d{2})\b", line):
                year = int(match.group(1))
                if year not in years:
                    years.append(year)
                if len(years) == 2:
                    return years
        return years

    @staticmethod
    def _financial_units(text: str) -> tuple[int, Decimal]:
        lowered = text.casefold()
        if "в миллионах" in lowered or re.search(r"\bмлн\.?\s*(?:руб|₽)", lowered):
            return 1_000_000, Decimal(1)
        if "в тысячах" in lowered or re.search(r"\bтыс\.?\s*(?:руб|₽)", lowered):
            return 1_000, Decimal(1)
        if "трлн" in lowered:
            # unit_scale is a 32-bit database column; store trillions as thousands of billions.
            return 1_000_000_000, Decimal(1_000)
        if "млрд" in lowered:
            return 1_000_000_000, Decimal(1)
        return 1, Decimal(1)

    @staticmethod
    def _primary_financial_page(text: str) -> bool:
        lowered = text.casefold()
        return any(
            marker in lowered
            for marker in (
                "консолидированный отчет о прибыли или убытке",
                "консолидированный отчёт о прибыли или убытке",
                "консолидированный отчет о финансовом положении",
                "консолидированный отчёт о финансовом положении",
                "пресс-релиз",
            )
        )

    @staticmethod
    def _decimal(value: str) -> Decimal | None:
        cleaned = value.strip().replace("\u00a0", "").replace(" ", "").replace("−", "-")
        negative = cleaned.startswith("(") and cleaned.endswith(")")
        cleaned = cleaned.strip("()").replace(",", ".")
        if cleaned.count(".") > 1:
            cleaned = cleaned.replace(".", "", cleaned.count(".") - 1)
        try:
            number = Decimal(cleaned)
            return -number if negative else number
        except InvalidOperation:
            return None
